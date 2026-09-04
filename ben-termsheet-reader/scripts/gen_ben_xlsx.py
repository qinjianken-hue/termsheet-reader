"""Build the BEN (Bonus Enhanced Note) tracker workbook from a JSON file.

Usage:
    python gen_ben_xlsx.py <data.json> [--outdir DIR] [--sender NAME]

Sheets
------
Tracker  Section 1 = the standard 10-column layout, ONE ROW PER UNDERLYING
         (worst-of basket), cells colour-coded:
             GREEN = cross-checked against the dealer place-order email / subject
             AMBER = termsheet only (read correctly, no independent second source)
             RED   = the email disagrees with the termsheet
         Section 2 = "Only for Rajiv" - ONE ROW PER NOTE, black header, no
         colour coding. Emitted ONLY when the sender is Rajiv (see IS_RAJIV).
Audit    One row per term: extracted value, the email reference it was checked
         against, match flag, and the source of the check.

See references/columns.md for the field dictionary and the JSON schema.
"""
import argparse
import datetime as dt
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---- palette (shared with the FCN reader) --------------------------------
GREEN = PatternFill("solid", fgColor="C6EFCE")
AMBER = PatternFill("solid", fgColor="FFEB9C")
RED = PatternFill("solid", fgColor="FFC7CE")
HDR = PatternFill("solid", fgColor="1F4E78")
BLACK = PatternFill("solid", fgColor="000000")
HDR_FONT = Font(color="FFFFFF", bold=True)

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BTHIN = Side(style="thin", color="000000")
BBORDER = Border(left=BTHIN, right=BTHIN, top=BTHIN, bottom=BTHIN)

# ---- layouts -------------------------------------------------------------
COLS = [
    "Trade Date", "Issue Date", "Tenor (mo)", "Final Valuation Date", "Underlying",
    "Currency", "Notional", "Coupon Flat", "Put Strike = Coupon Barrier", "Spot Price",
]
R_COLS = [
    "Order Date", "Product", "CCY", "Tenor", "Issue Date", "Final Valuation Date",
    "Maturity Date", "Underlying Name", "Put Strike %", "Coupon Flat %",
    "Coupon Barrier %", "Upside Participation %", "Size", "Strike Price",
]
# widths must serve both sections - they share columns A-N
WIDTHS = [13, 13, 10, 18, 13, 18, 14, 34, 26, 14, 16, 19, 14, 22]
WIDTHS_NO_RAJIV = [13, 13, 10, 18, 14, 10, 13, 13, 26, 13]


def is_rajiv(sender):
    """The 'Only for Rajiv' section is for this sender only."""
    return "rajiv" in (sender or "").lower()


def short_year(datestr):
    """'20-Aug-2026' -> '20-Aug-26'. String-only, so no locale surprises."""
    parts = (datestr or "").split("-")
    if len(parts) == 3 and len(parts[2]) == 4:
        return f"{parts[0]}-{parts[1]}-{parts[2][2:]}"
    return datestr


def strike_price(spot, put_strike_pct):
    return round(spot * put_strike_pct / 100.0, 4)


def fmt_price(x):
    """4 dp, trailing zeros stripped but never below 2 dp - matches how the
    prices are written in Rajiv's sheet (443.0555, 122.36, 182.00).
    Do NOT use %g here: it caps at 6 significant digits and silently drops the
    4th decimal on a 3-digit price (443.0555 -> 443.055)."""
    s = f"{x:.4f}"
    if "." in s:
        head, tail = s.split(".")
        tail = tail.rstrip("0").ljust(2, "0")
        s = f"{head}.{tail}"
    return s


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("data")
    ap.add_argument("--outdir", default=None)
    ap.add_argument("--sender", default=None)
    args = ap.parse_args()

    with open(args.data, encoding="utf-8") as fh:
        data = json.load(fh)

    notes = data.get("notes") or []
    if not notes:
        sys.exit("ERROR: no notes in the JSON.")
    sender = args.sender or data.get("sender") or ""
    rajiv = is_rajiv(sender)

    wb = Workbook()

    # ================= Tracker, section 1 =================================
    ws = wb.active
    ws.title = "Tracker"
    for c, name in enumerate(COLS, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill = HDR
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 30

    r = 2
    for n_i, note in enumerate(notes):
        if n_i:                      # blank row between different notes
            r += 1
        ver = note.get("verification", {})
        mismatched = note.get("mismatch", [])
        for u in note["underlyings"]:
            vals = [
                note["trade_date"], note["issue_date"], note["tenor"],
                note["final_val_date"], u["ticker"], note["currency"],
                note["notional"], note["coupon_flat"] / 100.0,
                note["put_strike"] / 100.0, u["spot"],
            ]
            for c, (name, v) in enumerate(zip(COLS, vals), 1):
                cell = ws.cell(row=r, column=c, value=v)
                if name in mismatched:
                    cell.fill = RED
                elif name in ver:
                    cell.fill = GREEN
                else:
                    cell.fill = AMBER
                cell.border = BORDER
                cell.alignment = Alignment(horizontal="center")
                if name in ("Coupon Flat", "Put Strike = Coupon Barrier"):
                    cell.number_format = "0.00%"
                elif name == "Notional":
                    cell.number_format = "#,##0"
                elif name == "Spot Price":
                    cell.number_format = "0.0000"
            r += 1

    # ================= Tracker, section 2: "Only for Rajiv" ===============
    if rajiv:
        title_row = ws.max_row + 3
        tc = ws.cell(row=title_row, column=1, value="Only for Rajiv")
        tc.font = Font(bold=True, size=12)

        hrow = title_row + 1
        for c, name in enumerate(R_COLS, 1):
            cell = ws.cell(row=hrow, column=c, value=name)
            cell.fill = BLACK
            cell.font = HDR_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BBORDER
        ws.row_dimensions[hrow].height = 30

        for i, note in enumerate(notes):
            names = "; ".join(u["name"] for u in note["underlyings"])
            strikes = "\n".join(
                f"{u.get('short', u['ticker'].split()[0])} : {note['currency']} "
                f"{fmt_price(strike_price(u['spot'], note['put_strike']))}"
                for u in note["underlyings"]
            )
            vals = [
                short_year(note["trade_date"]), note.get("product", "Regular BEN"),
                note["currency"], f"{note['tenor']}M",
                short_year(note["issue_date"]), short_year(note["final_val_date"]),
                short_year(note.get("maturity_date", "")), names,
                note["put_strike"], note["coupon_flat"],
                note.get("coupon_barrier", note["put_strike"]),
                note.get("upside_participation", 100),
                f"{note['currency']} {note['notional']:,}", strikes,
            ]
            rr = hrow + 1 + i
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=rr, column=c, value=v)
                cell.border = BBORDER
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            ws.row_dimensions[rr].height = 15 * max(2, len(note["underlyings"]))

    for i, w in enumerate(WIDTHS if rajiv else WIDTHS_NO_RAJIV, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ================= Audit ==============================================
    aw = wb.create_sheet("Audit")
    for c, name in enumerate(
            ["Term", "Extracted value", "Email reference", "Match", "Source of check"], 1):
        cell = aw.cell(row=1, column=c, value=name)
        cell.fill = HDR
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    row = 2
    for n_i, note in enumerate(notes):
        if len(notes) > 1:
            lc = aw.cell(row=row, column=1,
                         value=f"--- Note {n_i + 1}: {note.get('isin', '')} ---")
            lc.font = Font(bold=True)
            row += 1
        ver = note.get("verification", {})
        mismatched = note.get("mismatch", [])
        extracted = {
            "Trade Date": note["trade_date"],
            "Issue Date": note["issue_date"],
            "Tenor (mo)": str(note["tenor"]),
            "Final Valuation Date": note["final_val_date"],
            "Underlying": " / ".join(u["ticker"] for u in note["underlyings"]) + " (worst-of)",
            "Currency": note["currency"],
            "Notional": f"{note['notional']:,}",
            "Coupon Flat": f"{note['coupon_flat']}%",
            "Put Strike = Coupon Barrier": f"{note['put_strike']}%",
            "Spot Price": " / ".join(
                f"{u.get('short', u['ticker'].split()[0])} {fmt_price(u['spot'])}"
                for u in note["underlyings"]),
        }
        rows = [(name, extracted[name]) for name in COLS]
        rows += [(k, v) for k, v in (note.get("extra_audit") or {}).items()]

        for name, val in rows:
            if name in mismatched:
                ref = ver.get(name, {}).get("ref", "-")
                src = ver.get(name, {}).get("source", "")
                match, fill = "MISMATCH", RED
            elif name in ver:
                ref = ver[name].get("ref", "-")
                src = ver[name].get("source", "")
                match, fill = "OK", GREEN
            else:
                ref = "-"
                src = (note.get("amber_reason") or {}).get(name, "termsheet only")
                match, fill = "-", AMBER
            for c, v in enumerate([name, val, ref, match, src], 1):
                cell = aw.cell(row=row, column=c, value=v)
                cell.border = BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=(c == 5))
                if c == 4:
                    cell.fill = fill
                    cell.alignment = Alignment(horizontal="center")
            row += 1
        row += 1

    for i, w in enumerate([28, 42, 44, 10, 58], 1):
        aw.column_dimensions[get_column_letter(i)].width = w
    aw.freeze_panes = "A2"

    # ================= write ==============================================
    outdir = args.outdir or os.path.join(os.path.expanduser("~"), "Downloads",
                                         sender or "BEN")
    os.makedirs(outdir, exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y-%m-%d %H%M%S")
    name = f"BEN Termsheet Reader - {sender} {stamp}.xlsx" if sender \
        else f"BEN Termsheet Reader {stamp}.xlsx"
    out = os.path.join(outdir, name)
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
